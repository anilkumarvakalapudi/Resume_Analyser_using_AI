import pdfplumber
import docx
import io

MAX_FILE_SIZE = 3 * 1024 * 1024  # 3MB


# 🔹 MAIN FUNCTION (for local file usage)
def extract_text(file_path: str) -> str:
    with open(file_path, "rb") as f:
        file_bytes = f.read()

    _validate_file_size(file_bytes)

    file_path = file_path.lower()

    if file_path.endswith(".pdf"):
        text = _extract_pdf(file_bytes)

    elif file_path.endswith(".docx"):
        text = _extract_docx(file_bytes)

    else:
        raise ValueError("Unsupported file format. Use PDF or DOCX.")

    _validate_text(text)
    return text


# 🔹 FASTAPI VERSION (for UploadFile)
async def extract_text_from_upload(file) -> str:
    file_bytes = await file.read()

    _validate_file_size(file_bytes)

    filename = file.filename.lower()

    if filename.endswith(".pdf"):
        text = _extract_pdf(file_bytes)

    elif filename.endswith(".docx"):
        text = _extract_docx(file_bytes)

    else:
        raise ValueError("Unsupported file format. Use PDF or DOCX.")

    _validate_text(text)
    return text


# 🔹 FILE SIZE VALIDATION
def _validate_file_size(file_bytes):
    if len(file_bytes) > MAX_FILE_SIZE:
        raise ValueError("File exceeds 3MB limit")


# 🔹 TEXT VALIDATION
def _validate_text(text):
    if not text or not text.strip():
        raise ValueError(
            "No readable text found. File may be scanned or empty."
        )


# 🔹 PDF EXTRACTION
def _extract_pdf(file_bytes: bytes) -> str:
    text = ""

    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()

                if page_text:
                    text += page_text + "\n"

    except Exception as e:
        raise ValueError(f"Error reading PDF: {str(e)}")

    return text.strip()


# 🔹 DOCX EXTRACTION (PARAGRAPHS + TABLES)
def _extract_docx(file_bytes: bytes) -> str:
    try:
        doc = docx.Document(io.BytesIO(file_bytes))

        text = []

        # Extract paragraphs
        for p in doc.paragraphs:
            if p.text.strip():
                text.append(p.text.strip())

        # Extract tables (important for resumes)
        for table in doc.tables:
            for row in table.rows:
                row_text = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if row_text:
                    text.append(" | ".join(row_text))

        return "\n".join(text)

    except Exception as e:
        raise ValueError(f"Error reading DOCX: {str(e)}")